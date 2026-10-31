"""The admin's Excel file: sheet «Этапы» (the day's schedule) and sheet «Задания» (the task catalog).

Upload reads both sheets; download writes them back the same way, so the file can be edited and uploaded again.
"""
import io
import math
import re
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession

from app.models import Stage, Task
from app.stages import EVENT_TZ

STAGES_SHEET = "Этапы"
TASKS_SHEET = "Задания"

STAGE_COLUMNS = {
    "number": "Номер",
    "name": "Название",
    "starts_at": "Начало",
    "ends_at": "Конец",
    "message": "Сообщение при смене этапа",
}
TASK_COLUMNS = {
    "code": "ID",
    "stage": "Этап",
    "type": "Тип",
    "text": "Текст",
    "points": "Баллы",
    "coins": "Монеты",
    "partner_name": "Партнёр",
    "partner_package": "Пакет партнёра",
    "window_start": "Доступно с",
    "window_end": "Доступно до",
}
REQUIRED_STAGE = ("number", "name", "starts_at", "ends_at")
REQUIRED_TASK = ("code", "stage", "type", "text", "points")

# The 10 task types of TAG's rules (section 5). The file names them in Russian; the first name is what export writes.
TYPES = {
    "quiz": ["Квиз"],
    "selfie": ["Селфи"],
    "action": ["Действие"],
    "purchase": ["Покупка"],
    "registration": ["Регистрация"],
    "networking": ["Нетворкинг"],
    "crazy": ["Крейзи", "Крейзи-действие"],
    "attention": ["На внимательность"],
    "territory": ["Территориальное"],
    "pre": ["Предзадание"],
}
# Partner packages in feed order (TAG's rules, section 6)
PACKAGES = {
    "general": ["Генеральный партнёр", "General"],
    "headliner": ["Хедлайнер", "Хедлайнер форума", "Headliner"],
    "mf_partner": ["MF Партнёр", "MF Partner"],
    "supporters": ["Нас поддерживают", "Supporters"],
}
DATE_FORMATS = ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S")
EXCEL_DATE = "DD.MM.YYYY HH:MM"


def rule_coins(points: int, partner: bool) -> int:
    """Coins when the cell is blank: 1:1 with points; ×1.5 rounded up for partner tasks (TAG's rules, 3.1–3.2)."""
    return math.ceil(points * 1.5) if partner else points


@dataclass
class ImportResult:
    stages_added: int = 0
    stages_changed: int = 0
    added: int = 0
    changed: int = 0
    hidden: int = 0
    failed: list[dict] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)


class RowError(Exception):
    pass


def _lookup(names: dict[str, list[str]], raw: str) -> str | None:
    """«Действие ★», «Квиз (скрытый экспертный)», «Крейзи / на внимательность» → the first type named."""
    first = re.sub(r"\(.*?\)|★", "", raw).split("/")[0].strip().lower()
    for key, aliases in names.items():
        if first in (a.lower() for a in aliases):
            return key
    return None


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _int(v, reason: str) -> int:
    s = _text(v).replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        raise RowError(reason) from None
    if not n.is_integer() or n < 0:
        raise RowError(reason)
    return int(n)


def _time(v, reason: str) -> datetime:
    """A date-time cell, or text like 13.12.2026 09:00, in Dushanbe time."""
    if isinstance(v, datetime):
        dt = v
    else:
        s = _text(v)
        for fmt in DATE_FORMATS:
            try:
                dt = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        else:
            raise RowError(reason)
    return dt.replace(tzinfo=EVENT_TZ) if dt.tzinfo is None else dt


def _sheet_rows(ws, columns: dict[str, str]) -> tuple[dict[str, int], list[tuple[int, list]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, []
    header = [_text(h).lower() for h in rows[0]]
    cols = {key: header.index(name.lower()) for key, name in columns.items() if name.lower() in header}
    body = [(n, list(r)) for n, r in enumerate(rows[1:], start=2) if any(_text(c) for c in r)]
    return cols, body


def _cell(row: list, cols: dict[str, int], key: str):
    i = cols.get(key)
    return row[i] if i is not None and i < len(row) else None


def import_xlsx(db: DbSession, raw: bytes) -> ImportResult:
    result = ImportResult()
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        result.missing = [STAGES_SHEET, TASKS_SHEET]
        return result
    sheets = {name.strip().lower(): name for name in wb.sheetnames}
    result.missing = [s for s in (STAGES_SHEET, TASKS_SHEET) if s.lower() not in sheets]
    if result.missing:
        return result
    stage_cols, stage_rows = _sheet_rows(wb[sheets[STAGES_SHEET.lower()]], STAGE_COLUMNS)
    task_cols, task_rows = _sheet_rows(wb[sheets[TASKS_SHEET.lower()]], TASK_COLUMNS)
    result.missing = [f"{STAGES_SHEET}: {STAGE_COLUMNS[c]}" for c in REQUIRED_STAGE if c not in stage_cols] + [
        f"{TASKS_SHEET}: {TASK_COLUMNS[c]}" for c in REQUIRED_TASK if c not in task_cols]
    if result.missing:
        return result

    def fail(sheet: str, row: int, reason: str) -> None:
        result.failed.append({"sheet": sheet, "row": row, "reason": reason})

    stages = {s.number: s for s in db.scalars(select(Stage))}
    seen_stages: set[int] = set()
    for n, row in stage_rows:
        try:
            number = _int(_cell(row, stage_cols, "number"), "bad_stage_number")
            if number < 1:
                raise RowError("bad_stage_number")
            if number in seen_stages:
                raise RowError("duplicate_stage")
            name = _text(_cell(row, stage_cols, "name"))
            if not name:
                raise RowError("no_name")
            starts = _time(_cell(row, stage_cols, "starts_at"), "bad_start")
            ends = _time(_cell(row, stage_cols, "ends_at"), "bad_end")
            if ends <= starts:
                raise RowError("end_before_start")
        except RowError as e:
            fail(STAGES_SHEET, n, str(e))
            continue
        seen_stages.add(number)
        values = {"name": name, "starts_at": starts, "ends_at": ends,
                  "message": _text(_cell(row, stage_cols, "message"))}
        stage = stages.get(number)
        if stage is None:
            stage = Stage(number=number, **values)
            db.add(stage)
            stages[number] = stage
            result.stages_added += 1
        elif any(getattr(stage, k) != v for k, v in values.items()):
            for k, v in values.items():
                setattr(stage, k, v)
            result.stages_changed += 1
    db.flush()

    tasks = {t.code: t for t in db.scalars(select(Task))}
    seen: set[str] = set()
    for position, (n, row) in enumerate(task_rows):
        try:
            code = _text(_cell(row, task_cols, "code"))
            if not code:
                raise RowError("no_id")
            if code in seen:
                raise RowError("duplicate_id")
            stage_cell = _text(_cell(row, task_cols, "stage"))
            stage = stages.get(int(stage_cell)) if stage_cell.isdigit() else None
            if stage is None:
                raise RowError("unknown_stage")
            type_ = _lookup(TYPES, _text(_cell(row, task_cols, "type")))
            if type_ is None:
                raise RowError("unknown_type")
            text = _text(_cell(row, task_cols, "text"))
            if not text:
                raise RowError("no_text")
            if not _text(_cell(row, task_cols, "points")):
                raise RowError("no_points")
            points = _int(_cell(row, task_cols, "points"), "bad_points")
            partner_name = _text(_cell(row, task_cols, "partner_name"))
            package_cell = _text(_cell(row, task_cols, "partner_package"))
            package = _lookup(PACKAGES, package_cell) if package_cell else ""
            if package is None:
                raise RowError("unknown_package")
            if package and not partner_name:
                raise RowError("no_partner_name")
            coins_cell = _cell(row, task_cols, "coins")
            from_rule = not _text(coins_cell)
            coins = rule_coins(points, bool(partner_name)) if from_rule else _int(coins_cell, "bad_coins")
            w_start = _cell(row, task_cols, "window_start")
            w_end = _cell(row, task_cols, "window_end")
            window_start = _time(w_start, "bad_window") if _text(w_start) else None
            window_end = _time(w_end, "bad_window") if _text(w_end) else None
            if window_start and window_end and window_end <= window_start:
                raise RowError("bad_window")
        except RowError as e:
            fail(TASKS_SHEET, n, str(e))
            continue
        seen.add(code)
        values = {"stage_id": stage.id, "type": type_, "text": text, "points": points, "coins": coins,
                  "coins_from_rule": from_rule, "partner_name": partner_name, "partner_package": package,
                  "window_start": window_start, "window_end": window_end, "position": position, "active": True}
        task = tasks.get(code)
        if task is None:
            db.add(Task(code=code, **values))
            result.added += 1
        else:
            # Moving a row doesn't count as a change; editing it, or bringing a hidden one back, does
            if any(getattr(task, k) != v for k, v in values.items() if k != "position"):
                result.changed += 1
            for k, v in values.items():
                setattr(task, k, v)
    for code, task in tasks.items():
        if code not in seen and task.active:
            task.active = False
            result.hidden += 1
    db.commit()
    return result


def _local(dt: datetime | None) -> datetime | None:
    # Excel has no time zones: write Dushanbe wall-clock time, which is how the admin typed it
    return dt.astimezone(EVENT_TZ).replace(tzinfo=None) if dt else None


def export_xlsx(db: DbSession) -> bytes:
    """Both sheets with what is loaded now; with nothing loaded, an empty template with the right columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = STAGES_SHEET
    ws.append(list(STAGE_COLUMNS.values()))
    stages = list(db.scalars(select(Stage).order_by(Stage.number)))
    for s in stages:
        ws.append([s.number, s.name, _local(s.starts_at), _local(s.ends_at), s.message])
    for col in "CD":
        for cell in ws[col][1:]:
            cell.number_format = EXCEL_DATE

    wt = wb.create_sheet(TASKS_SHEET)
    wt.append(list(TASK_COLUMNS.values()))
    numbers = {s.id: s.number for s in stages}
    for t in db.scalars(select(Task).where(Task.active).order_by(Task.position, Task.id)):
        wt.append([t.code, numbers[t.stage_id], TYPES[t.type][0], t.text, t.points,
                   None if t.coins_from_rule else t.coins, t.partner_name,
                   PACKAGES[t.partner_package][0] if t.partner_package else None,
                   _local(t.window_start), _local(t.window_end)])
    for col in "IJ":
        for cell in wt[col][1:]:
            cell.number_format = EXCEL_DATE
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
