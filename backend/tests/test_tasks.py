import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import Workbook, load_workbook

from app import stages
from app.models import ScoreEvent
from app.db import SessionLocal
from tests.conftest import open_link, staff_headers

DUSHANBE = timezone(timedelta(hours=5))
STAGES = [
    [1, "Начало", "13.12.2026 09:00", "13.12.2026 10:30", "Игра началась!"],
    [2, "Первые победы", "13.12.2026 10:30", "13.12.2026 14:00", "Время побеждать"],
    [3, "Спад и разочарование", "13.12.2026 14:00", "13.12.2026 16:30", "Маркет открыт"],
    [4, "Последняя гонка", "13.12.2026 16:30", "13.12.2026 18:00", "Финишная прямая"],
]
STAGE_HEAD = ["Номер", "Название", "Начало", "Конец", "Сообщение при смене этапа"]
TASK_HEAD = ["ID", "Этап", "Тип", "Текст", "Баллы", "Монеты", "Партнёр", "Пакет партнёра", "Доступно с", "Доступно до"]


def task(code, stage, text, points=10, coins=None, type_="Действие", partner=None, package=None, start=None, end=None):
    return [code, stage, type_, text, points, coins, partner, package, start, end]


def catalog(n=43) -> list[list]:
    """n tasks spread over the 4 stages, like TAG's 43."""
    return [task(f"T{i:02d}", i % 4 + 1, f"Задание {i}", points=10 + i) for i in range(n)]


def xlsx(stage_rows=STAGES, task_rows=None, stage_head=STAGE_HEAD, task_head=TASK_HEAD) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Этапы"
    ws.append(stage_head)
    for r in stage_rows:
        ws.append(r)
    wt = wb.create_sheet("Задания")
    wt.append(task_head)
    for r in task_rows if task_rows is not None else catalog():
        wt.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(client, h, data: bytes):
    files = {"file": ("tasks.xlsx", data, "application/octet-stream")}
    return client.post("/api/admin/tasks/import", files=files, headers=h)


@pytest.fixture
def admin(client):
    return staff_headers(client, "admin")


@pytest.fixture
def guest(client, admin):
    from app.guest_list import import_csv
    with SessionLocal() as db:
        import_csv(db, "ID;Имя\nR-1;Иван\n".encode())
    with SessionLocal() as db:
        from app.models import Guest
        g = db.query(Guest).one()
    return open_link(client, f"http://api.test/b/{g.code}")


@pytest.fixture
def clock(monkeypatch):
    """Set the server clock to Dushanbe wall-clock time on the forum day."""
    def at(hhmm: str):
        h, m = map(int, hhmm.split(":"))
        t = datetime(2026, 12, 13, h, m, tzinfo=DUSHANBE)
        monkeypatch.setattr(stages, "now", lambda: t)
    return at


def feed(client, guest) -> dict:
    r = client.get("/api/feed", headers=guest)
    assert r.status_code == 200
    return r.json()


def test_upload_of_four_stages_and_43_tasks(client, admin):
    r = upload(client, admin, xlsx())
    assert r.status_code == 200
    body = r.json()
    assert (body["stages_added"], body["added"], body["failed"]) == (4, 43, [])


def test_bad_rows_fail_with_reason_and_others_load(client, admin):
    rows = [task("A", 1, "Хорошее"), task("B", 7, "Нет такого этапа"), task("C", 1, "Без баллов", points=None),
            task("D", 1, "Странный тип", type_="Танец"), task("A", 1, "Повтор ID")]
    body = upload(client, admin, xlsx(task_rows=rows)).json()
    assert body["added"] == 1
    assert body["failed"] == [
        {"sheet": "Задания", "row": 3, "reason": "unknown_stage"},
        {"sheet": "Задания", "row": 4, "reason": "no_points"},
        {"sheet": "Задания", "row": 5, "reason": "unknown_type"},
        {"sheet": "Задания", "row": 6, "reason": "duplicate_id"},
    ]


def test_file_without_a_sheet_or_column_says_what_is_missing(client, admin):
    wb = Workbook()
    wb.active.title = "Задания"
    out = io.BytesIO()
    wb.save(out)
    r = upload(client, admin, out.getvalue())
    assert r.status_code == 400
    assert r.json()["detail"] == {"code": "missing", "items": ["Этапы"]}
    r = upload(client, admin, xlsx(task_head=["ID", "Этап", "Тип", "Текст"]))
    assert r.json()["detail"]["items"] == ["Задания: Баллы"]


def test_blank_coins_follow_the_rule(client, admin, guest, clock):
    rows = [task("P", 1, "Партнёрское", points=15, partner="Coca-Cola", package="MF Партнёр"),
            task("N", 1, "Обычное", points=15), task("X", 1, "Свои монеты", points=15, coins=5)]
    upload(client, admin, xlsx(task_rows=rows))
    clock("09:30")
    coins = {t["text"]: (t["points"], t["coins"]) for t in feed(client, guest)["tasks"]}
    assert coins == {"Партнёрское": (15, 23), "Обычное": (15, 15), "Свои монеты": (15, 5)}


def test_reupload_changes_one_hides_one_and_leaves_the_rest(client, admin, guest, clock):
    upload(client, admin, xlsx())
    clock("09:30")
    before = {t["id"]: t for t in feed(client, guest)["tasks"]}
    rows = catalog()
    rows[0][3] = "Новый текст"  # T00, stage 1
    del rows[4]  # T04, stage 1
    body = upload(client, admin, xlsx(task_rows=rows)).json()
    assert (body["added"], body["changed"], body["hidden"], body["stages_changed"]) == (0, 1, 1, 0)
    after = {t["id"]: t for t in feed(client, guest)["tasks"]}
    gone = set(before) - set(after)
    assert len(gone) == 1 and before[gone.pop()]["text"] == "Задание 4"
    changed = [i for i in after if after[i] != before[i]]
    assert len(changed) == 1 and after[changed[0]]["text"] == "Новый текст"


def test_download_gives_back_what_is_loaded(client, admin):
    empty = load_workbook(io.BytesIO(client.get("/api/admin/tasks/export", headers=admin).content))
    assert empty.sheetnames == ["Этапы", "Задания"]
    assert [c.value for c in empty["Задания"][1]] == TASK_HEAD
    assert empty["Задания"].max_row == 1

    rows = catalog(3) + [task("P", 2, "Партнёр", points=15, partner="Coca-Cola", package="Хедлайнер",
                              start="13.12.2026 11:00", end="13.12.2026 11:05")]
    upload(client, admin, xlsx(task_rows=rows))
    data = client.get("/api/admin/tasks/export", headers=admin).content
    wb = load_workbook(io.BytesIO(data))
    assert wb["Этапы"]["C2"].value == datetime(2026, 12, 13, 9, 0)
    partner = [c.value for c in wb["Задания"][5]]
    assert partner[:8] == ["P", 2, "Действие", "Партнёр", 15, None, "Coca-Cola", "Хедлайнер"]
    assert partner[8:] == [datetime(2026, 12, 13, 11, 0), datetime(2026, 12, 13, 11, 5)]
    # Uploading the download again changes nothing
    body = upload(client, admin, data).json()
    assert (body["added"], body["changed"], body["hidden"], body["stages_changed"], body["failed"]) == (0, 0, 0, 0, [])


def test_feed_shows_only_the_current_stage_in_order(client, admin, guest, clock):
    rows = [
        task("plain", 2, "Обычное"),
        task("sup", 2, "Поддержка", partner="Б", package="Нас поддерживают"),
        task("head", 2, "Хедлайнер", partner="А", package="Хедлайнер"),
        task("quiz", 2, "Квиз", type_="Квиз", start="13.12.2026 11:00", end="13.12.2026 11:05"),
        task("closed", 2, "Квиз закрыт", type_="Квиз", start="13.12.2026 10:40", end="13.12.2026 10:45"),
        task("gen", 2, "Генеральный", partner="В", package="Генеральный партнёр"),
        task("other", 1, "Этап 1"),
    ]
    upload(client, admin, xlsx(task_rows=rows))
    clock("11:02")
    f = feed(client, guest)
    assert (f["state"], f["stage"]["number"], f["stage_count"]) == ("running", 2, 4)
    assert f["stage"]["ends_at"] == "2026-12-13T14:00:00+05:00"
    assert [t["text"] for t in f["tasks"]] == ["Квиз", "Генеральный", "Хедлайнер", "Поддержка", "Обычное"]
    assert f["tasks"][1]["partner_name"] == "В"


def test_stage_times_are_dushanbe_time(client, admin, guest, clock):
    upload(client, admin, xlsx(task_rows=[task("a", 1, "Утро")]))
    clock("08:59")
    assert feed(client, guest)["state"] == "before"
    clock("09:00")
    assert feed(client, guest)["state"] == "running"
    assert feed(client, guest)["starts_at"] == "2026-12-13T09:00:00+05:00"


def test_before_and_after_the_game(client, admin, guest, clock):
    upload(client, admin, xlsx())
    clock("07:00")
    f = feed(client, guest)
    assert (f["state"], f["stage"], f["tasks"]) == ("before", None, [])
    clock("18:00")
    assert feed(client, guest)["state"] == "over"


def test_next_stage_now(client, admin, guest, clock):
    upload(client, admin, xlsx())
    clock("11:00")
    r = client.post("/api/admin/stages/advance", headers=admin)
    assert r.json()["current"] == 3
    f = feed(client, guest)
    assert f["stage"]["number"] == 3
    assert f["stage"]["message"] == "Маркет открыт"
    assert {t["text"] for t in f["tasks"]} == {f"Задание {i}" for i in range(43) if i % 4 + 1 == 3}
    client.post("/api/admin/stages/advance", headers=admin)
    assert client.post("/api/admin/stages/advance", headers=admin).status_code == 409


def test_next_stage_before_the_game_starts_stage_one(client, admin, guest, clock):
    upload(client, admin, xlsx())
    clock("08:00")
    assert client.post("/api/admin/stages/advance", headers=admin).json()["current"] == 1


def test_task_card_details_and_done(client, admin, guest, clock):
    upload(client, admin, xlsx(task_rows=[task("a", 1, "Полный текст задания", points=25)]))
    clock("09:30")
    t = feed(client, guest)["tasks"][0]
    assert t["done"] is False
    r = client.get(f"/api/tasks/{t['id']}", headers=guest).json()
    assert (r["text"], r["points"], r["coins"]) == ("Полный текст задания", 25, 25)
    with SessionLocal() as db:
        db.add(ScoreEvent(guest_id=1, reason=f"task:{t['id']}", points=25))
        db.commit()
    assert feed(client, guest)["tasks"][0]["done"] is True
    assert client.get("/api/tasks/999", headers=guest).status_code == 404


def test_feed_needs_a_guest_and_admin_needs_the_password(client, admin):
    assert client.get("/api/feed").status_code == 401
    assert client.get("/api/feed", headers=admin).status_code == 401
    for method, path in [("get", "/api/admin/stages"), ("post", "/api/admin/stages/advance"),
                         ("post", "/api/admin/tasks/import"), ("get", "/api/admin/tasks/export")]:
        assert getattr(client, method)(path).status_code == 401
        assert getattr(client, method)(path, headers=staff_headers(client, "desk")).status_code == 401
